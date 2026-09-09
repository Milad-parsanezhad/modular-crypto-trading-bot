from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts" / "v17-strategy-lab"
OUT = ROOT / "deliverables" / "گزارش_نهایی_آزمایشگاه_استراتژی_ایچیموکو_v17.docx"
CHART = ART / "final_test_comparison.png"


SOURCES = [
    ("Deng et al.", "The profitability of Ichimoku Kinkohyo based trading rules in stock markets and FX markets", "International Journal of Finance & Economics, 2021", "https://doi.org/10.1002/ijfe.2067"),
    ("Che-Ngoc et al.", "Profitability of Ichimoku-Based Trading Rule in Vietnam Stock Market in the Context of the COVID-19 Outbreak", "Computational Economics, 2023", "https://doi.org/10.1007/s10614-022-10319-6"),
    ("Lutey and Rayome", "Ichimoku Cloud Forecasting Returns in the U.S.", "Global Business and Finance Review, 2022", "https://doi.org/10.17549/gbfr.2022.27.5.17"),
    ("Moskowitz, Ooi and Pedersen", "Time Series Momentum", "Journal of Financial Economics, 2012", "https://doi.org/10.1016/j.jfineco.2011.11.003"),
    ("Lempérière et al.", "Two Centuries of Trend Following", "Journal of Investment Strategies, 2014", "https://arxiv.org/abs/1404.3274"),
    ("Borgards and Czudaj", "Dynamic time series momentum of cryptocurrencies", "North American Journal of Economics and Finance, 2021", "https://doi.org/10.1016/j.najef.2021.101428"),
    ("Shen et al.", "Bitcoin intraday time-series momentum", "Financial Review 57(2), 2022", "https://doi.org/10.1111/fire.12290"),
    ("Lim, Zohren and Roberts", "Enhancing Time Series Momentum Strategies Using Deep Neural Networks", "Journal of Financial Data Science, 2019", "https://arxiv.org/abs/1904.04912"),
    ("Bailey et al.", "The Probability of Backtest Overfitting", "Journal of Computational Finance, 2016", "https://papers.ssrn.com/sol3/papers.cfm?abstract_id=2326253"),
    ("Bailey and López de Prado", "The Deflated Sharpe Ratio", "Journal of Portfolio Management, 2014", "https://papers.ssrn.com/sol3/papers.cfm?abstract_id=2460551"),
    ("Joubert", "Meta-Labeling Theory and Framework", "Journal of Financial Data Science, 2022", "https://doi.org/10.3905/jfds.2022.1.098"),
    ("Meyer, Barziy and Joubert", "Meta-Labeling Calibration and Position Sizing", "Journal of Financial Data Science, 2023", "https://doi.org/10.3905/jfds.2023.1.119"),
    ("World Cup Trading Championships", "Historical Standings", "Official competition records", "https://www.worldcupchampionships.com/world-cup-trading-championship-historical-standings"),
    ("Bysik and Ślepaczuk", "Machine Learning-Based Bitcoin Trading Under Transaction Costs", "Preprint, 2026", "https://arxiv.org/abs/2606.00060"),
]


def set_rtl(paragraph, rtl=True):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rpr.extend([color, underline]); run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t); hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_rtl(p)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 1
    set_repeat_header(table.rows[0])
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(value)
        set_cell_shading(cell, "1F4E79")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255); r.bold = True; r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cells[i], "EAF0F6")
            for p in cells[i].paragraphs:
                set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.font.size = Pt(8.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def percent(value):
    return f"{100 * float(value):.2f}%"


def make_chart(periods):
    final = periods[periods["period"].eq("final_test") & periods["strategy"].isin(["B0", "S1", "S3", "S4", "S6", "IRCP"])]
    pivot = final.pivot(index="strategy", columns="symbol", values="total_return").fillna(0) * 100
    ax = pivot.plot(kind="bar", figsize=(9, 4.8), color=["#1F4E79", "#7F8C8D"])
    ax.axhline(0, color="black", lw=0.8)
    ax.set_ylabel("Net return (%)"); ax.set_xlabel("")
    ax.set_title("Final test 2025 — net return after costs and risk limits")
    ax.grid(axis="y", alpha=0.2); plt.xticks(rotation=0)
    plt.tight_layout(); plt.savefig(CHART, dpi=200); plt.close()


def main():
    summary = json.loads((ART / "summary.json").read_text(encoding="utf-8"))
    periods = pd.read_csv(ART / "rule_strategy_period_summary.csv")
    costs = pd.read_csv(ART / "cost_sensitivity.csv")
    make_chart(periods)
    OUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "DejaVu Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "DejaVu Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "DejaVu Sans")
        style._element.rPr.rFonts.set(qn("w:cs"), "DejaVu Sans")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)

    title = doc.add_paragraph(style="Title")
    set_rtl(title); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("گزارش نهایی آزمایشگاه استراتژی ایچیموکو نسخه ۰٫۱۷")
    p = doc.add_paragraph(); set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ارزیابی علّی و هزینه‌واقعی روی بیت‌کوین و اتریوم در تایم‌فریم چهار ساعته").bold = True
    p = doc.add_paragraph(); set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("پروژه کارشناسی ارشد ربات معامله‌گر هوشمند رمزارز")
    doc.add_paragraph(); doc.add_paragraph()
    add_text(doc, "این گزارش پژوهش علمی، طراحی مهندسی و نتایج اجرای واقعی نسخه ۰٫۱۷ را یکپارچه می‌کند. نتیجه اصلی آن است که IRGC-S در شکل فعلی رد می‌شود، در حالی که شکست روند S6 در آزمون نهایی سال ۲۰۲۵ روی هر دو دارایی بازده خالص مثبت ایجاد کرده است. این نتیجه تنها مجوز ادامه پژوهش و Paper Trading است و مجوز معامله زنده نیست.")
    doc.add_page_break()

    heading(doc, "۱ خلاصه مدیریتی")
    add_text(doc, "داده رسمی Binance Vision شامل ۱۳٬۱۵۱ کندل چهار ساعته برای هر یک از BTC/USDT و ETH/USDT از ابتدای ۲۰۲۰ تا پایان ۲۰۲۵ بود. تصمیم‌ها در پایان کندل تولید و در قیمت بازشدن کندل بعدی اجرا شدند. هزینه پایه برای هر سمت معامله ۱۰ واحد پایه کارمزد و ۲ واحد پایه لغزش بود؛ بنابراین اصطکاک کامل رفت‌وبرگشت ۲۴ واحد پایه محاسبه شد.")
    add_text(doc, "در آزمون نهایی ۲۰۲۵، S6 روی BTC بازده ۳٫۵۳٪ با Sharpe برابر ۰٫۸۳ و روی ETH بازده ۵٫۰۶٪ با Sharpe برابر ۱٫۲۱ ثبت کرد. IRGC-S از ۱۳۹ رویداد خارج‌نمونه، ۴۹ معامله را برگزید اما بازده آن منفی ۰٫۴۶٪، نرخ برد ۴۶٫۹٪ و Profit Factor برابر ۰٫۹۴ بود. در نتیجه IRGC-S ارتقا نیافت.")
    add_text(doc, "بهترین نتیجه سالانه یا یک زیر‌دوره برای ادعای سودآوری کافی نیست. S6 باید در اجرای forward paper، داده صرافی دوم، هزینه‌های بدبینانه و تحلیل ظرفیت نیز پایدار بماند. مدل‌های پیچیده مانند LSTM، Transformer یا RL تا عبور این baseline ساده از دروازه‌های سخت‌گیرانه نباید جایگزین آن شوند.")

    heading(doc, "۲ شواهد علمی و منطق طراحی")
    heading(doc, "۲ ۱ ایچیموکو و ناپایداری زمانی", 2)
    add_text(doc, "مطالعه Deng و همکاران سودآوری قواعد ایچیموکو را در چند بازار سهام و ارز بررسی کرد و نشان داد عملکرد در زیر‌دوره‌ها پایدار نیست؛ پارامترهای پیش‌فرض نباید حقیقت بهینه تلقی شوند [۱]. مطالعه بازار ویتنام عملکرد بهتر را در شرایط خاص همه‌گیری گزارش کرد، اما همین وابستگی به دوره و بازار، تعمیم جهان‌شمول را محدود می‌کند [۲]. مطالعه Lutey و Rayome نیز بیشتر شواهد پیش‌بینی‌پذیری آماری محدود ارائه می‌کند و نه تضمین سود قابل‌معامله پس از هزینه [۳].")
    heading(doc, "۲ ۲ روند و شکست", 2)
    add_text(doc, "ادبیات time-series momentum نشان می‌دهد روند در طبقات مختلف دارایی مشاهده شده است [۴،۵]. بااین‌حال قدرت روند به افق، بازار و هزینه حساس است. پژوهش‌های رمزارز نیز وجود مومنتوم را در برخی نمونه‌ها تأیید کرده‌اند، ولی نتایج کوتاه‌افق و نوسان‌مدیریت‌شده یکسان نیست [۶،۷]. این شواهد استفاده از S6 را به‌عنوان فرضیه قابل‌آزمون توجیه می‌کند، نه به‌عنوان قانون قطعی.")
    heading(doc, "۲ ۳ یادگیری ماشین و کنترل هزینه", 2)
    add_text(doc, "مدل‌های عمیق می‌توانند در برخی مجموعه‌های آتی، روند و اندازه موقعیت را هم‌زمان بیاموزند، اما مزیت گزارش‌شده به هزینه و طراحی تابع هدف وابسته است [۸]. پژوهش تازه روی بیت‌کوین نیز نشان می‌دهد مدل‌هایی با عملکرد ناخالص مثبت ممکن است با هزینه ده واحد پایه شکست بخورند و فیلتر هزینه نقش تعیین‌کننده داشته باشد [۱۴]. بنابراین معیار اصلی پروژه بازده خالص خارج‌نمونه است، نه Accuracy یا AUC به‌تنهایی.")
    heading(doc, "۲ ۴ بیش‌برازش و مسابقات معامله‌گری", 2)
    add_text(doc, "بازده‌های بسیار بزرگ مسابقات برای تولید فرضیه مفیدند، اما به علت انتخاب برنده، افق کوتاه، اهرم و نبود افشای کامل روش، شاهد مستقیم آلفای پایدار نیستند [۱۳]. مطالعات PBO و Deflated Sharpe توضیح می‌دهند که انتخاب بهترین نتیجه از میان آزمایش‌های متعدد می‌تواند یک راهبرد فاقد مهارت را برنده نشان دهد [۹،۱۰]. ازاین‌رو پارامترها از پیش ثبت و نتایج منفی حفظ شده‌اند.")

    doc.add_page_break()
    heading(doc, "۳ معماری و راهبردهای منجمد")
    rows = [
        ("B0", "خرید و نگهداری", "معیار اقتصادی بدون Kill Switch"),
        ("S1", "کراس تنکان و کیجون", "ورود کراس صعودی و خروج کراس نزولی"),
        ("S2", "شکست کومو", "عبور قیمت از ابر و خروج زیر کیجون"),
        ("S3", "رژیم ایچیموکو", "قیمت بالای ابر، تنکان بالای کیجون و شیب کیجون مثبت"),
        ("S4", "پولبک تأییدشده", "لمس ناحیه تنکان کیجون در رژیم صعودی"),
        ("S5", "رد حمایت ساختاری", "واکنش صعودی به کیجون یا لبه ابر"),
        ("S6", "شکست ۲۰ کندلی", "شکست سقف قبلی با شیب مثبت EMA200"),
        ("IRCP", "پولبک رژیم‌آگاه", "S4 همراه با تأیید علّی چیکو"),
        ("C1", "بازگشت به میانگین", "سه افت متوالی در روند بلندمدت صعودی"),
        ("C2", "شکست فشردگی", "فشردگی Bollinger و شکست با حجم"),
        ("IRGC-S", "رویداد و Meta Label", "رژیم، CUSUM، سیگنال اولیه، Triple Barrier و Logistic"),
    ]
    add_table(doc, ["شناسه", "خانواده", "تعریف اجرایی"], rows, [2.0, 4.5, 9.5])
    add_text(doc, "ویژگی‌های ایچیموکو بدون انتقال آینده ساخته شدند. Senkou صرفاً با داده قابل‌مشاهده در همان زمان محاسبه شد و چیکو به مقایسه قیمت جاری با قیمت ۲۶ کندل قبل تبدیل گردید. سطوح شکست نیز با shift یک کندل، کندل جاری را از تعریف مقاومت حذف می‌کنند.")

    heading(doc, "۴ داده و روش آزمایش")
    add_table(doc, ["مولفه", "مقدار"], [
        ("منبع", "آرشیو رسمی ماهانه Binance Vision"),
        ("بازار", "BTC/USDT و ETH/USDT"),
        ("دوره", "۲۰۲۰-۰۱-۰۱ تا ۲۰۲۵-۱۲-۳۱"),
        ("تعداد", "۱۳٬۱۵۱ کندل برای هر دارایی"),
        ("توسعه", "۲۰۲۰ تا ۲۰۲۳"),
        ("اعتبارسنجی", "۲۰۲۴"),
        ("آزمون نهایی", "۲۰۲۵"),
        ("هزینه پایه", "۱۲ واحد پایه یک‌طرفه؛ ۲۴ واحد پایه رفت‌وبرگشت"),
        ("ریسک", "۰٫۲۵٪ بودجه معامله و سقف وزن ۳۵٪ برای هر دارایی"),
        ("Kill Switch", "افت ۵٪؛ توقف تا پایان اجرای پژوهش و نیاز به بازنشانی بیرونی"),
    ], [5.0, 11.0])
    add_text(doc, "تمام ردیف‌های warm-up حفظ شدند، ولی امتیازدهی هر دوره فقط در بازه همان دوره انجام گرفت. B0 عمداً از Kill Switch عبور نکرد تا خرید و نگهداری واقعی باقی بماند. راهبردهای فعال با اندازه موقعیت مبتنی بر ATR و بودجه ریسک ارزیابی شدند.")

    heading(doc, "۵ نتایج تجربی")
    leaders = []
    for symbol in ("BTC/USDT", "ETH/USDT"):
        for period in ("development", "validation", "final_test"):
            group = periods[periods["symbol"].eq(symbol) & periods["period"].eq(period)]
            r = group.sort_values("sharpe", ascending=False).iloc[0]
            leaders.append((symbol, period, r["strategy"], percent(r["total_return"]), f"{r['sharpe']:.2f}", percent(r["max_drawdown"])))
    add_table(doc, ["دارایی", "دوره", "رهبر Sharpe", "بازده خالص", "Sharpe", "MDD"], leaders, [2.8, 3.0, 2.5, 3.0, 2.0, 2.7])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(CHART), width=Inches(6.2))
    cap = doc.add_paragraph("شکل ۱ مقایسه بازده خالص راهبردهای منتخب در آزمون نهایی ۲۰۲۵")
    set_rtl(cap); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading(doc, "۵ ۱ تحلیل S6", 2)
    add_text(doc, "S6 تنها راهبردی بود که در آزمون نهایی ۲۰۲۵ روی هر دو دارایی بازده مثبت و Sharpe مثبت معنادار از نظر اقتصادی اولیه داشت. بااین‌حال بازده آن محدود بود و هنوز آزمون آماری مستقل، داده صرافی دوم و Paper Trading لازم است. برتری S6 نسبت به B0 در سال ۲۰۲۵ از نظر کنترل افت سرمایه چشمگیر است، اما انتخاب آن پس از مشاهده مجموعه‌ای از قواعد، خطر multiple testing ایجاد می‌کند.")
    heading(doc, "۵ ۲ تحلیل IRCP", 2)
    btc_ircp = periods[(periods.symbol.eq("BTC/USDT")) & periods.strategy.eq("IRCP")]
    rows = [(r.period, percent(r.total_return), f"{r.sharpe:.2f}", percent(r.max_drawdown)) for r in btc_ircp.itertuples()]
    add_table(doc, ["دوره BTC", "بازده", "Sharpe", "MDD"], rows, [4.0, 4.0, 3.5, 4.0])
    add_text(doc, "IRCP در توسعه و اعتبارسنجی BTC مطلوب بود، اما در آزمون نهایی ۲۰۲۵ رهبر نشد. در ETH نیز پایداری لازم را نشان نداد. این شکست انتقال بین دارایی‌ها با ادبیات ناپایداری ایچیموکو سازگار است و مانع تبدیل IRCP به موتور انحصاری ربات می‌شود.")
    heading(doc, "۵ ۳ تحلیل IRGC-S", 2)
    irgc = summary["irgc_s"]
    m = irgc["metrics_selected"]
    add_table(doc, ["شاخص", "نتیجه"], [
        ("کل رویدادها", irgc["events_total"]),
        ("رویدادهای OOS", irgc["events_oos"]),
        ("معاملات انتخاب‌شده", irgc["events_selected"]),
        ("fold مثبت", irgc["positive_oos_folds"]),
        ("بازده خالص", percent(m["total_return"])),
        ("نرخ برد", percent(m["win_rate"])),
        ("Profit Factor", f"{m['profit_factor']:.2f}"),
        ("تصمیم", irgc["decision"]),
    ], [7.0, 9.0])
    add_text(doc, "Meta-labeling زیان سیگنال اولیه را کاهش داد، اما آن را مثبت نکرد. کاهش زیان از ۶٫۷۲٪ در سیگنال اولیه به ۰٫۴۶٪ در نمونه انتخاب‌شده، بدون foldهای مثبت پایدار، برای ارتقا کافی نیست. این نتیجه نشان می‌دهد فیلتر ML می‌تواند تعداد معاملات را کم کند، ولی کاهش معامله مترادف با ایجاد آلفا نیست.")

    heading(doc, "۶ حساسیت هزینه و محدودیت ریسک")
    stress_rows = []
    for (symbol, cost), group in costs.groupby(["symbol", "one_way_cost_bps"]):
        group = group[~group["strategy"].eq("B0")]
        row = group.sort_values("sharpe", ascending=False).iloc[0]
        stress_rows.append((symbol, f"{2*cost:.0f} bps", row.strategy, percent(row.total_return), f"{row.sharpe:.2f}"))
    add_table(doc, ["دارایی", "هزینه رفت‌وبرگشت", "بهترین راهبرد", "بازده کل ۲۰۲۰–۲۰۲۵", "Sharpe"], stress_rows, [3.0, 4.0, 3.0, 3.0, 3.0])
    add_text(doc, "محدودیت‌های فعلی شامل سقف وزن ۳۵٪ برای هر دارایی، سقف ناخالص ۷۰٪ برای پرتفوی دو دارایی، بودجه ریسک ۰٫۲۵٪ و Kill Switch پنج‌درصدی است. اجرای زنده در قرارداد نرم‌افزاری خاموش باقی مانده و هیچ مدل یا سیگنالی اجازه دورزدن موتور مستقل ریسک را ندارد.")

    heading(doc, "۷ تصمیم پژوهشی و ادامه ساخت ربات")
    add_text(doc, "نسخه v0.17 از نظر مهندسی Strategy Lab را تکمیل می‌کند، اما از نظر علمی فقط S6 را برای مرحله بعد نامزد می‌کند. IRGC-S و مدل ML فعلی رد شده‌اند و نباید در اجرای واقعی سرمایه به کار روند. معماری ربات باید ensemble باقی بماند، ولی فعال‌سازی هر عضو به عبور مستقل از دروازه خارج‌نمونه وابسته است.")
    add_text(doc, "مرحله بعدی مجاز عبارت است از forward paper trading برای S6 و IRCP به‌عنوان challenger، ثبت کامل سفارش فرضی، لغزش، تأخیر، خطای داده و تطبیق موجودی. سپس باید بازآزمایی روی CoinEx یا منبع مستقل، Bootstrap اختلاف بازده، تصحیح multiple testing و تحلیل ظرفیت انجام شود. تنها پس از عبور این مراحل می‌توان Testnet را بررسی کرد؛ Live همچنان ممنوع است.")

    heading(doc, "۸ محدودیت‌ها")
    for text in (
        "داده‌ها متعلق به یک منبع آرشیوی‌اند و خطا یا تفاوت کیفیت اجرای صرافی‌های دیگر را پوشش نمی‌دهند.",
        "بک‌تست کندلی ترتیب دقیق رخدادهای داخل کندل را نمی‌بیند؛ در برخورد هم‌زمان حد سود و ضرر، سناریوی بدبینانه انتخاب شده است.",
        "انتخاب بهترین راهبرد از بین چند نامزد خطر data snooping دارد و به آزمون آماری اصلاح‌شده نیاز دارد.",
        "هزینه ثابت جایگزین کامل spread متغیر، market impact، partial fill و latency واقعی نیست.",
        "نتیجه مثبت S6 در سال ۲۰۲۵ کوتاه است و تضمین تکرار در آینده محسوب نمی‌شود.",
        "ویژگی‌های funding، open interest و on-chain عمداً از آزمایش اصلی حذف شدند تا اثر قاعده قیمت قابل شناسایی بماند.",
    ):
        p = doc.add_paragraph(style="List Bullet"); set_rtl(p); p.add_run(text)

    heading(doc, "۹ قابلیت بازتولید")
    add_text(doc, "کد در مخزن modular-crypto-trading-bot قرار می‌گیرد. ماژول اصلی research_bot/strategy_lab.py، بارگذار آرشیو research_bot/binance_spot_archive.py، آداپتر امن research_bot/forward_candidate_v17.py، اجرای آزمایش scripts/run_strategy_lab_v17.py و مجموعه تست‌های v0.17 هستند. فایل summary.json تصمیم ماشینی، داده مورد استفاده، پیکربندی و وضعیت اجرای زنده را ثبت می‌کند.")

    heading(doc, "منابع")
    for i, (author, title_text, venue, url) in enumerate(SOURCES, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(f"{i}. {author}. {title_text}. {venue}. ")
        add_hyperlink(p, url, url)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("گزارش پژوهشی v0.17  |  اجرای زنده غیرفعال")
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
